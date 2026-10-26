"""Document -> per-page text (native text and OCR) and cheap regex candidates for profile fields.

Page-level decisions
--------------------
Every page is judged on its own: a PDF page with a text layer is used as is, a page without one (or with a
garbled one) is rendered and OCR'd, so mixed native/scanned PDFs work. A DOCX that is only pictures is
OCR'd image by image. A page whose OCR could not be completed (engine down, output cut off by the token
limit) is reported as *incomplete* instead of silently becoming empty text, so the caller can retry it
rather than conclude that the resume says nothing.

OCR engines
-----------
Qwen vision first, Tesseract second. Each sits behind a circuit breaker: after N consecutive failures the
engine is skipped for a cooldown and then probed again, so one transient error no longer switches OCR off
for the rest of a multi-day run.
"""

from __future__ import annotations

import base64
import hashlib
import io
import logging
import os
import re
import threading
import time
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from PIL import Image

from profile_enrichment.config import CONFIG

logger = logging.getLogger(__name__)

EXTRACTOR_VERSION = "2026-10-26.v2"
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def normalize_text(value: Optional[str]) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def has_meaningful_text(text: Optional[str], minimum_chars: Optional[int] = None) -> bool:
    content = normalize_text(text)
    alnum_chars = sum(1 for ch in content if ch.isalnum())
    return alnum_chars >= (minimum_chars or CONFIG["ocr_min_text_chars"])


def _alnum_count(text: str) -> int:
    return sum(1 for ch in text if ch.isalnum())


def _looks_garbled(text: str) -> bool:
    """A text layer that is mostly replacement characters / CID placeholders is not usable text."""
    if not text.strip():
        return False
    bad = text.count("\ufffd") + 5 * text.count("(cid:")
    visible = sum(1 for ch in text if not ch.isspace())
    return visible > 0 and (bad / visible > 0.2 or _alnum_count(text) / visible < 0.3)


# ---------------------------------------------------------------------------------------------
# document model
# ---------------------------------------------------------------------------------------------

@dataclass
class PageText:
    number: int
    text: str = ""
    source: str = "native"  # native | ocr_vision | ocr_tesseract | none
    ocr_attempted: bool = False
    ocr_complete: bool = True
    error: str = ""


@dataclass
class DocumentContent:
    pages: List[PageText] = field(default_factory=list)
    extension: str = ""
    sha256: str = ""
    problems: List[str] = field(default_factory=list)  # whole-document problems (corrupt, encrypted, ...)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def page_texts(self) -> List[Tuple[int, str]]:
        return [(p.number, p.text) for p in self.pages if p.text.strip()]

    @property
    def ocr_complete(self) -> bool:
        return all(p.ocr_complete for p in self.pages)

    @property
    def incomplete_pages(self) -> List[int]:
        return [p.number for p in self.pages if not p.ocr_complete]

    @property
    def native_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.source == "native" and p.text.strip())

    @property
    def ocr_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.source.startswith("ocr") and p.text.strip())

    def failure_reason(self) -> str:
        """Why this document cannot be trusted as the complete resume ('' when it can)."""
        if self.problems:
            return "; ".join(self.problems)
        if not self.ocr_complete:
            details = "; ".join(f"page {p.number}: {p.error or 'ocr incomplete'}" for p in self.pages if not p.ocr_complete)
            return f"ocr_incomplete ({details})"
        return ""


# ---------------------------------------------------------------------------------------------
# circuit breaker + OCR engines
# ---------------------------------------------------------------------------------------------

class CircuitBreaker:
    """closed -> (N consecutive failures) -> open for `cooldown` seconds -> half-open (one probe) -> closed."""

    def __init__(self, name: str, threshold: int, cooldown: float, clock=time.monotonic):
        self.name, self.threshold, self.cooldown, self._clock = name, threshold, cooldown, clock
        self._lock = threading.Lock()
        self._failures = 0
        self._opened_at: Optional[float] = None
        self._probing = False
        self.last_error = ""

    def allow(self) -> bool:
        with self._lock:
            if self._opened_at is None:
                return True
            if self._clock() - self._opened_at >= self.cooldown and not self._probing:
                self._probing = True  # exactly one caller probes the engine again
                return True
            return False

    def success(self) -> None:
        with self._lock:
            if self._opened_at is not None:
                logger.info("OCR_ENGINE_RECOVERED | engine=%s", self.name)
            self._failures, self._opened_at, self._probing = 0, None, False

    def failure(self, reason: str) -> None:
        with self._lock:
            self._failures += 1
            self.last_error = reason
            if self._probing or self._failures >= self.threshold:
                if self._opened_at is None or self._probing:
                    logger.warning("OCR_ENGINE_PAUSED | engine=%s | cooldown=%ss | reason=%s", self.name, self.cooldown, reason)
                self._opened_at, self._probing = self._clock(), False

    @property
    def is_open(self) -> bool:
        with self._lock:
            return self._opened_at is not None


_BREAKERS: Dict[str, CircuitBreaker] = {}
_BREAKER_LOCK = threading.Lock()


def breaker(name: str) -> CircuitBreaker:
    with _BREAKER_LOCK:
        if name not in _BREAKERS:
            _BREAKERS[name] = CircuitBreaker(name, CONFIG["vision_ocr_failure_threshold"], CONFIG["vision_ocr_cooldown"])
        return _BREAKERS[name]


def reset_breakers() -> None:
    with _BREAKER_LOCK:
        _BREAKERS.clear()


@dataclass
class OcrOutcome:
    text: str = ""
    complete: bool = False
    engine: str = ""
    error: str = ""


def _vision_base_url() -> Optional[str]:
    return CONFIG["vision_ocr_base_url"] or CONFIG["local_model_base_url"]


def _vision_configured() -> bool:
    return bool(CONFIG["vision_ocr_enabled"] and _vision_base_url() and CONFIG["vision_ocr_model"])


def _tesseract_configured() -> bool:
    return bool(CONFIG["ocr_enabled"])


def _resolve_tesseract_cmd(configured: str) -> str:
    """Accept either the tesseract executable or the folder that contains it."""
    if configured and os.path.isdir(configured):
        for candidate in ("tesseract.exe", "tesseract"):
            path = os.path.join(configured, candidate)
            if os.path.isfile(path):
                return path
    return configured or ""


def _image_to_data_uri(image: Image.Image) -> str:
    rgb = image.convert("RGB")
    max_side = CONFIG["vision_ocr_max_image_side"]
    if max_side and max(rgb.size) > max_side:
        ratio = max_side / max(rgb.size)
        rgb = rgb.resize((max(1, int(rgb.width * ratio)), max(1, int(rgb.height * ratio))))
    buffer = io.BytesIO()
    rgb.save(buffer, format="PNG", optimize=True)
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def _vision_call(image: Image.Image) -> Tuple[str, str]:
    """One vision request -> (text, finish_reason). Raises on transport errors."""
    from profile_enrichment.inference import _get_openai_client

    client = _get_openai_client(
        _vision_base_url(), CONFIG["vision_ocr_api_key"] or CONFIG["local_model_api_key"] or "local-qwen",
        CONFIG["vision_ocr_timeout"],
    )
    response = client.chat.completions.create(
        model=CONFIG["vision_ocr_model"],
        messages=[
            {"role": "system", "content": "You are an OCR engine. Transcribe the page verbatim as plain text. No commentary."},
            {"role": "user", "content": [
                {"type": "text", "text": (
                    "Transcribe every piece of text in this resume page, preserving reading order. Include names, "
                    "phone numbers, emails, addresses and headings. If the page has no text, reply with an empty response.")},
                {"type": "image_url", "image_url": {"url": _image_to_data_uri(image)}},
            ]},
        ],
        temperature=0,
        max_tokens=CONFIG["vision_ocr_max_tokens"],
    )
    choice = response.choices[0]
    return choice.message.content or "", choice.finish_reason or "stop"


def _vision_ocr(image: Image.Image, depth: int = 0) -> OcrOutcome:
    """Vision OCR with retries. Output cut off by the token limit is not accepted: the page is split and re-read."""
    last_error = ""
    for attempt in (1, 2):
        try:
            text, finish = _vision_call(image)
        except Exception as exc:  # noqa: BLE001
            last_error = f"{type(exc).__name__}: {exc}"
            if "not found" in last_error.lower() or "NotFoundError" in last_error:
                break  # a missing model will not appear on a retry
            time.sleep(0.5 * attempt)
            continue
        if finish == "length":
            if depth >= 1 or image.height < 200:
                return OcrOutcome(text, False, "vision", "ocr output truncated at max_tokens")
            half = image.height // 2
            top = _vision_ocr(image.crop((0, 0, image.width, min(image.height, half + 40))), depth + 1)
            bottom = _vision_ocr(image.crop((0, max(0, half - 40), image.width, image.height)), depth + 1)
            return OcrOutcome(f"{top.text}\n{bottom.text}".strip(), top.complete and bottom.complete, "vision",
                              top.error or bottom.error)
        return OcrOutcome(text, True, "vision")
    return OcrOutcome("", False, "vision", last_error)


def _tesseract_ocr(image: Image.Image) -> OcrOutcome:
    try:
        import pytesseract

        resolved = _resolve_tesseract_cmd(CONFIG["tesseract_cmd"])
        if resolved:
            pytesseract.pytesseract.tesseract_cmd = resolved
        return OcrOutcome(pytesseract.image_to_string(image.convert("RGB"), timeout=90), True, "tesseract")
    except Exception as exc:  # noqa: BLE001
        return OcrOutcome("", False, "tesseract", f"{type(exc).__name__}: {exc}")


_ocr_clock = threading.local()


def consume_ocr_seconds() -> float:
    """Seconds this thread spent in OCR since the last call (lets the pipeline time OCR apart from parsing)."""
    seconds = getattr(_ocr_clock, "seconds", 0.0)
    _ocr_clock.seconds = 0.0
    return seconds


def ocr_image(image: Image.Image) -> OcrOutcome:
    """OCR one page image through the first engine whose breaker allows it. Never raises."""
    started = time.monotonic()
    try:
        return _ocr_image(image)
    finally:
        _ocr_clock.seconds = getattr(_ocr_clock, "seconds", 0.0) + time.monotonic() - started


def _ocr_image(image: Image.Image) -> OcrOutcome:
    errors: List[str] = []
    for name, configured, run in (("vision", _vision_configured(), _vision_ocr), ("tesseract", _tesseract_configured(), _tesseract_ocr)):
        if not configured:
            continue
        gate = breaker(name)
        if not gate.allow():
            errors.append(f"{name}: paused ({gate.last_error})")
            continue
        outcome = run(image)
        if outcome.complete:
            gate.success()
            return outcome
        gate.failure(outcome.error or "incomplete output")
        errors.append(f"{name}: {outcome.error}")
    return OcrOutcome("", False, "", "; ".join(errors) or "no OCR engine configured")


def ocr_engine_status() -> Dict[str, str]:
    status = {}
    status["vision"] = ("not configured" if not _vision_configured()
                        else ("paused: " + breaker("vision").last_error) if breaker("vision").is_open else "active")
    status["tesseract"] = ("not configured" if not _tesseract_configured()
                           else ("paused: " + breaker("tesseract").last_error) if breaker("tesseract").is_open else "active")
    return status


def check_vision_ready() -> Tuple[bool, str]:
    """Prove the *vision* model can read: OCR a rendered line of text and compare. Independent of the text model."""
    if not _vision_configured():
        return False, "vision OCR not configured (set VISION_OCR_BASE_URL / VISION_OCR_MODEL)"
    from PIL import ImageDraw, ImageFont

    image = Image.new("RGB", (900, 160), "white")
    try:
        font = ImageFont.load_default(size=48)
    except TypeError:  # very old Pillow
        font = ImageFont.load_default()
    ImageDraw.Draw(image).text((20, 50), "Cardiology 4827", fill="black", font=font)
    outcome = _vision_ocr(image)
    if not outcome.complete:
        return False, f"vision model call failed: {outcome.error}"
    if "4827" not in outcome.text.replace(" ", ""):
        return False, f"vision model answered but did not read the test image (got {outcome.text[:60]!r})"
    return True, f"vision model {CONFIG['vision_ocr_model']} read the test image"


def check_tesseract_ready() -> Tuple[bool, str]:
    if not _tesseract_configured():
        return False, "OCR_ENABLED is false"
    image = Image.new("RGB", (600, 120), "white")
    from PIL import ImageDraw, ImageFont

    try:
        font = ImageFont.load_default(size=48)
    except TypeError:
        font = ImageFont.load_default()
    ImageDraw.Draw(image).text((10, 30), "Nurse 7351", fill="black", font=font)
    outcome = _tesseract_ocr(image)
    if not outcome.complete:
        return False, outcome.error
    return ("7351" in outcome.text.replace(" ", "")), f"tesseract returned {outcome.text.strip()[:40]!r}"


# ---------------------------------------------------------------------------------------------
# parsers
# ---------------------------------------------------------------------------------------------

def _ocr_page(page: PageText, image: Image.Image) -> None:
    page.ocr_attempted = True
    outcome = ocr_image(image)
    page.ocr_complete = outcome.complete
    page.error = outcome.error
    if outcome.complete:
        page.text = outcome.text
        page.source = "ocr_vision" if outcome.engine == "vision" else "ocr_tesseract"
    elif not page.text.strip():
        page.source = "none"


def _page_needs_ocr(native: str) -> bool:
    if not CONFIG["ocr_skip_when_native_text"]:
        return True
    return _alnum_count(native) < CONFIG["ocr_min_page_chars"] or _looks_garbled(native)


def _parse_pdf(data: bytes, ocr_enabled: bool) -> DocumentContent:
    from pypdf import PdfReader

    doc = DocumentContent(extension=".pdf")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            doc.problems.append("pdf_encrypted")
            return doc
        natives = []
        for page in reader.pages:
            try:
                natives.append(page.extract_text() or "")
            except Exception as exc:  # noqa: BLE001
                logger.debug("pdf page text extraction failed: %s", exc)
                natives.append("")
    except Exception as exc:  # noqa: BLE001
        doc.problems.append(f"pdf_unreadable: {type(exc).__name__}: {exc}")
        return doc

    doc.pages = [PageText(number=i + 1, text=text) for i, text in enumerate(natives)]
    needing = [p for p in doc.pages if _page_needs_ocr(p.text)]
    if not needing:
        return doc
    if not ocr_enabled:
        # OCR was switched off on purpose: report what we have, do not claim the document was read in full.
        for page in needing:
            if not page.text.strip():
                page.ocr_complete = False
                page.error = "page has no text layer and OCR is disabled"
                page.source = "none"
        return doc
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        for page in needing:
            page.ocr_complete, page.error = False, f"cannot render page for OCR: {type(exc).__name__}: {exc}"
        return doc
    try:
        for page in needing:
            try:
                image = pdf[page.number - 1].render(scale=2).to_pil()
            except Exception as exc:  # noqa: BLE001
                page.ocr_complete, page.error = False, f"render failed: {type(exc).__name__}: {exc}"
                continue
            native = page.text
            _ocr_page(page, image)
            if not page.ocr_complete and native.strip():
                page.text, page.source, page.ocr_complete = native, "native", True  # garbled-but-present beats nothing
    finally:
        pdf.close()
    return doc


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_paragraph_pages(data: bytes) -> List[str]:
    """Text of every paragraph in document order (body, tables, text boxes, headers/footers), split on page breaks."""
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(io.BytesIO(data))
    pages: List[List[str]] = [[]]
    seen_in_textbox: set = set()

    def paragraph_text(p) -> str:
        return "".join(t.text or "" for t in p.iter(qn("w:t")))

    def has_break(p) -> bool:
        if any(True for _ in p.iter(qn("w:lastRenderedPageBreak"))):
            return True
        return any(br.get(qn("w:type")) == "page" for br in p.iter(qn("w:br")))

    # headers/footers first: names and contact lines often live there
    header_texts: List[str] = []
    for section in document.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            try:
                for paragraph in part.paragraphs:
                    text = paragraph.text.strip()
                    if text and text not in header_texts:
                        header_texts.append(text)
            except Exception:  # noqa: BLE001
                continue
    pages[0].extend(header_texts)

    for p in document.element.body.iter(qn("w:p")):
        in_fallback = any(anc.tag.endswith("}Fallback") for anc in p.iterancestors())
        if in_fallback:
            continue  # text boxes appear twice (Choice + Fallback)
        text = paragraph_text(p).strip()
        in_box = any(anc.tag == qn("w:txbxContent") for anc in p.iterancestors())
        if text:
            key = (text, in_box)
            if in_box and key in seen_in_textbox:
                continue
            seen_in_textbox.add(key)
            in_cell = any(anc.tag == qn("w:tc") for anc in p.iterancestors())
            pages[-1].append(text if not in_cell else text)
        if has_break(p):
            pages.append([])
    return ["\n".join(lines) for lines in pages if lines]


def _docx_images(data: bytes) -> List[Image.Image]:
    images: List[Image.Image] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith("word/media/"))
        for name in names:
            raw = archive.read(name)
            if len(raw) < 4000:
                continue  # bullets, rules, logos
            try:
                image = Image.open(io.BytesIO(raw))
                image.load()
            except Exception:  # noqa: BLE001 - emf/wmf/svg and other non-raster media
                continue
            if image.width < 150 or image.height < 150:
                continue
            images.append(image)
    return images


def _parse_docx(data: bytes, ocr_enabled: bool) -> DocumentContent:
    doc = DocumentContent(extension=".docx")
    try:
        texts = _docx_paragraph_pages(data)
    except Exception as exc:  # noqa: BLE001
        doc.problems.append(f"docx_unreadable: {type(exc).__name__}: {exc}")
        return doc
    doc.pages = [PageText(number=i + 1, text=t) for i, t in enumerate(texts)]
    native_chars = sum(_alnum_count(p.text) for p in doc.pages)
    if native_chars >= CONFIG["ocr_min_text_chars"] and CONFIG["ocr_skip_when_native_text"]:
        return doc

    try:
        images = _docx_images(data)
    except Exception as exc:  # noqa: BLE001
        doc.problems.append(f"docx_media_unreadable: {type(exc).__name__}: {exc}")
        return doc
    if not images:
        if native_chars == 0:
            doc.pages = doc.pages or []
        return doc
    if not ocr_enabled:
        if native_chars == 0:
            doc.pages.append(PageText(number=1, source="none", ocr_complete=False, error="image-only document and OCR is disabled"))
        return doc
    base = len(doc.pages)
    for offset, image in enumerate(images, start=1):
        page = PageText(number=base + offset, source="none")
        _ocr_page(page, image)
        doc.pages.append(page)
    return doc


def _rtf_to_text(raw: str) -> str:
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "ignore"), raw)
    text = re.sub(r"\\(?:par|line)\b", "\n", text)
    text = re.sub(r"\{\\\*[^{}]*\}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    return re.sub(r"[{}]", "", text)


def _parse_image(data: bytes, ocr_enabled: bool) -> DocumentContent:
    doc = DocumentContent(extension=".img")
    try:
        image = Image.open(io.BytesIO(data))
        frames = []
        for index in range(getattr(image, "n_frames", 1)):
            image.seek(index)
            frames.append(image.copy())
    except Exception as exc:  # noqa: BLE001
        doc.problems.append(f"image_unreadable: {type(exc).__name__}: {exc}")
        return doc
    for index, frame in enumerate(frames, start=1):
        page = PageText(number=index, source="none")
        if ocr_enabled:
            _ocr_page(page, frame)
        else:
            page.ocr_complete, page.error = False, "image-only document and OCR is disabled"
        doc.pages.append(page)
    return doc


def sniff_extension(data: bytes, file_ref: str = "") -> str:
    """The document type from its bytes (URLs sometimes lie or have no extension)."""
    if data[:5] == b"%PDF-":
        return ".pdf"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return ".docx"
        except zipfile.BadZipFile:
            pass
    if data[:4] == bytes([0xD0, 0xCF, 0x11, 0xE0]):
        return ".doc"
    if data[:5] == b"{" + bytes([0x5C]) + b"rtf":
        return ".rtf"
    ext = os.path.splitext(file_ref.split("?")[0].lower())[1]
    return ext or ".bin"


def extract_document(data: bytes, file_extension: str, ocr_enabled: bool | None = None) -> DocumentContent:
    """Per-page content of a document. Raises ValueError for an unsupported type; never for a bad file."""
    ext = file_extension.lower()
    if not ext.startswith("."):
        ext = f".{ext}"
    ocr_enabled = CONFIG["ocr_enabled"] if ocr_enabled is None else ocr_enabled
    if ext == ".pdf":
        doc = _parse_pdf(data, ocr_enabled)
    elif ext == ".docx":
        doc = _parse_docx(data, ocr_enabled)
    elif ext == ".txt":
        doc = DocumentContent(pages=[PageText(1, data.decode("utf-8", errors="ignore"))])
    elif ext == ".rtf":
        doc = DocumentContent(pages=[PageText(1, _rtf_to_text(data.decode("latin-1", errors="ignore")))])
    elif ext in _IMAGE_EXTENSIONS:
        doc = _parse_image(data, ocr_enabled)
    else:
        raise ValueError(f"Unsupported document type: {ext}")
    doc.extension = ext
    doc.sha256 = hashlib.sha256(data).hexdigest()
    return doc


def extract_text_layers_from_bytes(data: bytes, file_extension: str, ocr_enabled: bool | None = None) -> Dict[str, str]:
    """Back-compatible view of :func:`extract_document` (native/ocr/combined text)."""
    doc = extract_document(data, file_extension, ocr_enabled)
    return {"native_text": doc.native_text, "ocr_text": doc.ocr_text, "combined_text": doc.text}


def extract_text_layers_from_file(path: str) -> Dict[str, str]:
    with open(path, "rb") as handle:
        return extract_text_layers_from_bytes(handle.read(), os.path.splitext(path)[1])


# ---------------------------------------------------------------------------------------------
# regex candidates (cheap, high precision for contact details)
# ---------------------------------------------------------------------------------------------

_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA", "colorado": "CO",
    "connecticut": "CT", "delaware": "DE", "district of columbia": "DC", "florida": "FL", "georgia": "GA",
    "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
    "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD", "massachusetts": "MA",
    "michigan": "MI", "minnesota": "MN", "mississippi": "MS", "missouri": "MO", "montana": "MT",
    "nebraska": "NE", "nevada": "NV", "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM",
    "new york": "NY", "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
    "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}
_STATE_CODES = set(_STATES.values()) | {"PR"}
HEADER_CHARS = 700  # contact details live at the top; a state name deep in the body is not the candidate's location


def _state_from_text(text: str) -> Optional[str]:
    cleaned = normalize_text(text)
    if not cleaned:
        return None
    match = re.search(r",\s*([A-Z]{2})(?:\b|$)", cleaned)
    if match and match.group(1) in _STATE_CODES:
        return match.group(1)
    lowered = cleaned.lower()
    for state_name, code in sorted(_STATES.items(), key=lambda kv: -len(kv[0])):
        if re.search(r"\b" + re.escape(state_name) + r"\b", lowered):
            return code
    return None


_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_PHONE_RES = (
    re.compile(r"(?<!\d)(?:\+?1[\s.\-]*)?\(?([2-9]\d{2})\)?[\s.\-]*([2-9]\d{2})[\s.\-]*(\d{4})(?!\d)"),
)


def find_emails(text: str) -> List[Tuple[str, int]]:
    """Every email in the text with its offset, original capitalisation kept."""
    return [(m.group(0).rstrip(".-"), m.start()) for m in _EMAIL_RE.finditer(text)]


def find_phones(text: str) -> List[Tuple[str, int]]:
    """Every US-style phone number with its offset (year ranges such as 2020-2024 never match)."""
    found: List[Tuple[str, int]] = []
    for pattern in _PHONE_RES:
        for m in pattern.finditer(text):
            found.append((normalize_text(m.group(0)), m.start()))
    return sorted(found, key=lambda item: item[1])


def _extract_phone(text: str) -> Optional[str]:
    phones = find_phones(text)
    return phones[0][0] if phones else None


def extract_profile_fields_from_text(text: str) -> Dict[str, Any]:
    extracted: Dict[str, Any] = {}
    cleaned = text.replace("\r", "\n")

    def safe_get(pattern: str, flags: int = 0, group: int = 1) -> Optional[str]:
        match = re.search(pattern, cleaned, flags=flags)
        return normalize_text(match.group(group)) if match else None

    emails = find_emails(cleaned)
    if emails:
        extracted["email"] = emails[0][0]
    phone = _extract_phone(cleaned)
    if phone:
        extracted["phone"] = phone

    name_match = re.search(r"(?im)^\s*([A-Z][A-Za-z'`-]+(?:\s+[A-Z][A-Za-z'`-]+){1,3})\s*$", cleaned)
    if name_match:
        name_parts = name_match.group(1).strip().split()
        if len(name_parts) >= 2:
            extracted["first_name"] = name_parts[0]
            extracted["last_name"] = name_parts[-1]

    location = safe_get(r"(?im)^\s*(?:Location|Address|Based in)\b\s*[:\-]\s*([A-Za-z0-9 .,'#-]+?)\s*$")
    if location:
        extracted["location"] = location
        if "," in location:
            extracted["city"] = normalize_text(location.split(",", 1)[0])
    city = safe_get(r"(?im)^\s*City\b\s*[:\-]\s*([A-Za-z .'-]+?)\s*$")
    if city:
        extracted["city"] = city.split(",", 1)[0].strip()
        extracted["location"] = extracted.get("location") or city

    # "City, ST" in the header block (first lines only) - the common resume layout
    header = cleaned[:HEADER_CHARS]
    header_loc = re.search(r"(?m)^[^\n@\d]{0,40}?([A-Z][A-Za-z.' -]{1,30}),\s*([A-Z]{2})\b", header)
    if header_loc and header_loc.group(2) in _STATE_CODES and "city" not in extracted:
        extracted["city"] = normalize_text(header_loc.group(1))
        extracted["state_code"] = header_loc.group(2)
        extracted["state"] = header_loc.group(2)
    state_code = _state_from_text(extracted.get("location") or city or header)
    if state_code:
        extracted["state_code"] = state_code
        extracted["state"] = state_code

    specialty = safe_get(
        r"(?im)^\s*(?:Specialty|Specialties|Focus|Practice Area|Profession Type|Profession)\b"
        r"\s*[:\-]\s*([A-Za-z /&.()-]+?)\s*$"
    )
    if specialty:
        extracted["specialty"] = specialty
    return extracted


# ---------------------------------------------------------------------------------------------
# legacy AI entry points (the pipeline now calls profile_enrichment.inference directly)
# ---------------------------------------------------------------------------------------------

_legacy_state = threading.local()


def consume_ai_call_failure() -> str:
    """Legacy: error text of the last :func:`run_ai_extraction` call on this thread (use ExtractionResult instead)."""
    error = getattr(_legacy_state, "error", "")
    _legacy_state.error = ""
    return error


def run_ai_extraction(document_path: str, text: str) -> Dict[str, Any]:
    """Legacy dict view over :func:`inference.extract_fields`; a failure sets consume_ai_call_failure()."""
    del document_path
    _legacy_state.error = ""
    if not CONFIG["ai_extractor_enabled"] or not text.strip():
        return {}
    from profile_enrichment import inference

    result = inference.extract_fields([(1, text)])
    if not result.succeeded:
        _legacy_state.error = f"{result.status.value}: {result.error}"
        return {}
    return {c.field: c.value for c in result.candidates}


def combine_extracted_values(base: Dict[str, Any], extra: Dict[str, Any]) -> Dict[str, Any]:
    merged = dict(base)
    for key, value in extra.items():
        if value in (None, "", " "):
            continue
        if key not in merged or not merged[key] or not str(merged[key]).strip():
            merged[key] = value
    return merged
