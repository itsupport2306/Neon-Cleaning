"""Per-page OCR decisions, mixed native/scanned PDFs, image-only DOCX, truncated OCR and the recoverable breaker."""

import io

import pytest
from PIL import Image, ImageDraw

from profile_enrichment import extractors
from profile_enrichment.extractors import CircuitBreaker, OcrOutcome, extract_document


def make_pdf(page_texts):
    """Minimal valid PDF; a page with text=None is a blank (image-like) page without a text layer."""
    objects = []

    def add(body):
        objects.append(body)
        return len(objects)

    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    pages_id_placeholder = len(objects) + 1 + 2 * len(page_texts)
    for text in page_texts:
        stream = "" if text is None else f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET"
        content = add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        page = add(f"<< /Type /Page /Parent {pages_id_placeholder} 0 R /MediaBox [0 0 612 792] "
                   f"/Contents {content} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>")
        page_ids.append(page)
    pages = add("<< /Type /Pages /Kids [" + " ".join(f"{i} 0 R" for i in page_ids) + f"] /Count {len(page_ids)} >>")
    assert pages == pages_id_placeholder
    catalog = add(f"<< /Type /Catalog /Pages {pages} 0 R >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def picture(text="SCANNED RESUME", size=(420, 220)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    for row in range(0, size[1], 6):  # noise so the PNG exceeds the "tiny logo" size filter
        draw.line([(0, row), (size[0], (row * 7) % size[1])], fill=(200 + row % 50, 210, 220))
    draw.text((20, 20), text, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def fresh_ocr(monkeypatch):
    extractors.reset_breakers()
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", False)  # engines are replaced per test
    monkeypatch.setitem(extractors.CONFIG, "ocr_enabled", True)
    yield
    extractors.reset_breakers()


def fake_ocr(monkeypatch, text="Scanned page text. Registered Nurse. ICU.", complete=True, error=""):
    calls = []

    def _ocr(image):
        calls.append(image.size)
        return OcrOutcome(text if complete else "", complete, "tesseract", error)

    monkeypatch.setattr(extractors, "ocr_image", _ocr)
    return calls


NATIVE = "Jordan Lee RN ICU Registered Nurse with many years of bedside experience and strong clinical skills"


def test_pdf_with_text_layer_needs_no_ocr(monkeypatch):
    calls = fake_ocr(monkeypatch)
    doc = extract_document(make_pdf([NATIVE]), ".pdf")
    assert calls == [] and doc.pages[0].source == "native" and "Jordan Lee" in doc.text and doc.ocr_complete


def test_mixed_pdf_ocrs_only_the_scanned_pages(monkeypatch):
    calls = fake_ocr(monkeypatch)
    doc = extract_document(make_pdf([NATIVE, None, NATIVE.replace("Jordan", "Casey")]), ".pdf")
    assert [p.source for p in doc.pages] == ["native", "ocr_tesseract", "native"]
    assert len(calls) == 1  # per page, not per document: the two text pages are not re-read
    assert doc.ocr_complete and "Scanned page text" in doc.text and "Casey Lee" in doc.text


def test_scanned_page_with_failing_ocr_is_incomplete_not_empty(monkeypatch):
    fake_ocr(monkeypatch, complete=False, error="engine down")
    doc = extract_document(make_pdf([NATIVE, None]), ".pdf")
    assert not doc.ocr_complete and doc.incomplete_pages == [2]
    assert "ocr_incomplete" in doc.failure_reason() and "engine down" in doc.failure_reason()


def test_ocr_disabled_does_not_pretend_the_scanned_page_was_read(monkeypatch):
    calls = fake_ocr(monkeypatch)
    doc = extract_document(make_pdf([NATIVE, None]), ".pdf", ocr_enabled=False)
    assert calls == [] and not doc.ocr_complete and "OCR is disabled" in doc.failure_reason()


def test_image_only_docx_is_ocrd_image_by_image(monkeypatch):
    from docx import Document

    calls = fake_ocr(monkeypatch)
    document = Document()
    for label in ("SCANNED RESUME ONE", "SCANNED RESUME TWO"):  # identical images would be stored once by Word
        document.add_picture(io.BytesIO(picture(label)))
    buffer = io.BytesIO()
    document.save(buffer)

    doc = extract_document(buffer.getvalue(), ".docx")
    assert len(calls) == 2 and doc.ocr_complete and doc.text.count("Scanned page text") == 2


def test_docx_with_text_does_not_trigger_ocr(monkeypatch):
    from docx import Document

    calls = fake_ocr(monkeypatch)
    document = Document()
    document.add_paragraph(NATIVE)
    buffer = io.BytesIO()
    document.save(buffer)
    doc = extract_document(buffer.getvalue(), ".docx")
    assert calls == [] and NATIVE in doc.text


def test_docx_header_footer_and_page_breaks_are_read(monkeypatch):
    from docx import Document
    from docx.enum.text import WD_BREAK

    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header: Jordan Lee, RN, Austin TX"
    document.add_paragraph("Page one body " * 12)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Page two body " * 12)
    buffer = io.BytesIO()
    document.save(buffer)
    doc = extract_document(buffer.getvalue(), ".docx", ocr_enabled=False)
    assert "Jordan Lee" in doc.text and len([p for p in doc.pages if p.text.strip()]) == 2


def test_corrupt_documents_are_reported_not_raised():
    assert "pdf_unreadable" in extract_document(b"%PDF-1.4 garbage", ".pdf").failure_reason()
    assert "docx_unreadable" in extract_document(b"PK\x03\x04nope", ".docx").failure_reason()
    with pytest.raises(ValueError):
        extract_document(b"x", ".xyz")


# ---- OCR engines ------------------------------------------------------------------------------

def test_circuit_breaker_opens_then_recovers_after_the_cooldown():
    now = [0.0]
    cb = CircuitBreaker("vision", threshold=2, cooldown=60, clock=lambda: now[0])
    assert cb.allow()
    cb.failure("boom")
    assert cb.allow()
    cb.failure("boom")
    assert cb.is_open and not cb.allow()  # paused, not permanently disabled
    now[0] = 61
    assert cb.allow() and not cb.allow()  # exactly one probe after the cooldown
    cb.failure("still down")
    assert not cb.allow()
    now[0] = 130
    assert cb.allow()
    cb.success()
    assert not cb.is_open and cb.allow()


def test_one_transient_vision_error_does_not_disable_ocr_for_the_run(monkeypatch):
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", True)
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_base_url", "http://127.0.0.1:1")
    monkeypatch.setitem(extractors.CONFIG, "ocr_enabled", False)  # vision only
    outcomes = iter([OcrOutcome("", False, "vision", "timeout"), OcrOutcome("page text", True, "vision")])
    monkeypatch.setattr(extractors, "_vision_ocr", lambda image: next(outcomes))
    image = Image.new("RGB", (50, 50))
    assert not extractors.ocr_image(image).complete
    second = extractors.ocr_image(image)
    assert second.complete and second.text == "page text"


def test_vision_failure_falls_back_to_tesseract_for_that_page(monkeypatch):
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", True)
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_base_url", "http://127.0.0.1:1")
    monkeypatch.setattr(extractors, "_vision_ocr", lambda image: OcrOutcome("", False, "vision", "404 model not found"))
    monkeypatch.setattr(extractors, "_tesseract_ocr", lambda image: OcrOutcome("tesseract text", True, "tesseract"))
    outcome = extractors.ocr_image(Image.new("RGB", (50, 50)))
    assert outcome.complete and outcome.engine == "tesseract"


def test_truncated_vision_output_is_not_accepted_as_a_complete_page(monkeypatch):
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", True)
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_base_url", "http://127.0.0.1:1")
    calls = []

    def call(image):
        calls.append(image.height)
        return ("cut off mid sen", "length") if image.height > 500 else ("half of the page", "stop")

    monkeypatch.setattr(extractors, "_vision_call", call)
    outcome = extractors._vision_ocr(Image.new("RGB", (600, 800)))
    assert outcome.complete and "half of the page" in outcome.text  # re-read as two halves
    assert len(calls) == 3

    monkeypatch.setattr(extractors, "_vision_call", lambda image: ("still cut", "length"))
    assert not extractors._vision_ocr(Image.new("RGB", (600, 800))).complete


def test_missing_vision_model_is_reported_by_its_own_check(monkeypatch):
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", True)
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_base_url", "http://127.0.0.1:1")
    monkeypatch.setattr(extractors, "_vision_call", lambda image: (_ for _ in ()).throw(RuntimeError("NotFoundError: model not found")))
    ok, detail = extractors.check_vision_ready()
    assert not ok and "failed" in detail
    monkeypatch.setitem(extractors.CONFIG, "vision_ocr_enabled", False)
    assert extractors.check_vision_ready()[0] is False


def test_sniffing_does_not_trust_the_url_extension():
    assert extractors.sniff_extension(make_pdf([NATIVE]), "https://x/y.bin") == ".pdf"
    assert extractors.sniff_extension(b"zzzz", "https://x/y.docx?sig=1") == ".docx"
