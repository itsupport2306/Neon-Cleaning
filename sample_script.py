"""
Resume extractor: downloads resumes and uses an LLM (via Ollama) to extract
name, email, phone, specialty. No regex extraction.

Setup (Windows PowerShell):
    pip install requests pymupdf python-docx ollama
    ollama pull qwen2.5:7b        # text resumes
    ollama pull qwen2.5vl:7b      # scanned resumes (use 3b if low VRAM)

Input CSV  (resumes.csv) needs columns:  id,resume_url
    (resume_url can also be a local file path)

Run:
    python extract_resumes.py resumes.csv results.jsonl

It is resumable: rerun the same command and finished ids are skipped.
"""
import csv
import io
import json
import os
import sys
import tempfile
import time
from pathlib import Path

import fitz  # PyMuPDF
import ollama
import requests
from docx import Document

TEXT_MODEL = "qwen2.5:7b"
VISION_MODEL = "qwen2.5vl:7b"
MIN_TEXT_CHARS = 200      # below this, treat the file as a scan
MAX_TEXT_CHARS = 12000    # truncate long resumes
MAX_SCAN_PAGES = 2        # pages sent to the vision model
RETRIES = 2

PROMPT = """You are extracting contact details from a resume / CV.
Return ONLY a JSON object with exactly these keys:
{"name": string or null, "email": string or null, "phone": string or null, "specialty": string or null}

Rules:
- name: the candidate's full name (usually at the top).
- email: the candidate's email address exactly as written.
- phone: the candidate's phone number exactly as written, including country code if shown.
- specialty: the candidate's main professional or medical specialty / field
  (e.g. "Cardiology", "Pediatrics", "Software Engineering"). Use a short standard term.
- If a value is not present, use null. Do not guess or invent values.
"""


def download(url: str, dest_dir: str) -> str:
    if os.path.exists(url):
        return url
    r = requests.get(url, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    ctype = r.headers.get("content-type", "").lower()
    ext = Path(url.split("?")[0]).suffix.lower()
    if ext not in (".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg"):
        if "pdf" in ctype:
            ext = ".pdf"
        elif "wordprocessingml" in ctype:
            ext = ".docx"
        elif "png" in ctype:
            ext = ".png"
        elif "jpeg" in ctype or "jpg" in ctype:
            ext = ".jpg"
        else:
            ext = ".pdf"
    path = os.path.join(dest_dir, f"resume{ext}")
    with open(path, "wb") as f:
        f.write(r.content)
    return path


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        with fitz.open(path) as doc:
            return "\n".join(page.get_text() for page in doc).strip()
    if ext == ".docx":
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    return ""  # images and .doc -> treated as scans / unsupported


def render_images(path: str, dest_dir: str) -> list:
    ext = Path(path).suffix.lower()
    if ext in (".png", ".jpg", ".jpeg"):
        return [path]
    if ext == ".pdf":
        out = []
        with fitz.open(path) as doc:
            for i, page in enumerate(doc):
                if i >= MAX_SCAN_PAGES:
                    break
                p = os.path.join(dest_dir, f"page{i + 1}.png")
                page.get_pixmap(dpi=150).save(p)
                out.append(p)
        return out
    return []


def parse_json(raw: str) -> dict:
    raw = raw.strip()
    if raw.startswith("```"):
        raw = raw.strip("`")
        raw = raw[raw.find("{"):]
    start, end = raw.find("{"), raw.rfind("}")
    data = json.loads(raw[start:end + 1])
    keys = ("name", "email", "phone", "specialty")
    clean = {}
    for k in keys:
        v = data.get(k)
        if isinstance(v, str):
            v = v.strip()
            if v.lower() in ("", "null", "none", "n/a", "unknown"):
                v = None
        clean[k] = v
    return clean


def ask_text_model(text: str) -> dict:
    resp = ollama.chat(
        model=TEXT_MODEL,
        format="json",
        options={"temperature": 0},
        messages=[{"role": "user", "content": f"{PROMPT}\n\nRESUME TEXT:\n{text[:MAX_TEXT_CHARS]}"}],
    )
    return parse_json(resp["message"]["content"])


def ask_vision_model(images: list) -> dict:
    resp = ollama.chat(
        model=VISION_MODEL,
        format="json",
        options={"temperature": 0},
        messages=[{"role": "user", "content": PROMPT, "images": images}],
    )
    return parse_json(resp["message"]["content"])


def process(resume_id: str, url: str) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        path = download(url, tmp)
        text = extract_text(path)
        if len(text) >= MIN_TEXT_CHARS:
            result = ask_text_model(text)
            result["method"] = "text"
        else:
            images = render_images(path, tmp)
            if not images:
                raise ValueError("no extractable text and no renderable pages")
            result = ask_vision_model(images)
            result["method"] = "vision"
    return result


def main():
    if len(sys.argv) < 3:
        print("Usage: python extract_resumes.py resumes.csv results.jsonl")
        sys.exit(1)
    in_csv, out_jsonl = sys.argv[1], sys.argv[2]

    done = set()
    if os.path.exists(out_jsonl):
        with open(out_jsonl, encoding="utf-8") as f:
            for line in f:
                try:
                    rec = json.loads(line)
                    if rec.get("error") is None:
                        done.add(str(rec["id"]))
                except Exception:
                    pass

    with open(in_csv, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    todo = [r for r in rows if str(r["id"]) not in done]
    print(f"{len(rows)} rows, {len(done)} already done, {len(todo)} to process")

    with open(out_jsonl, "a", encoding="utf-8") as out:
        for n, row in enumerate(todo, 1):
            rid, url = str(row["id"]), row["resume_url"]
            rec = {"id": rid, "error": None}
            for attempt in range(RETRIES + 1):
                try:
                    rec.update(process(rid, url))
                    rec["error"] = None
                    break
                except Exception as e:
                    rec["error"] = f"{type(e).__name__}: {e}"
                    time.sleep(2)
            out.write(json.dumps(rec, ensure_ascii=False) + "\n")
            out.flush()
            status = rec["error"] or f'{rec.get("method")} -> {rec.get("specialty")}'
            print(f"[{n}/{len(todo)}] id={rid} {status}")


if __name__ == "__main__":
    main()